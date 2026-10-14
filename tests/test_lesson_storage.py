import copy
import json
import os
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from config.colleges import get_college_config
from services.lesson_plan_service import generate_lesson_plan, PendingLessonPlan
from services.supabase_service import SupabaseService, SupabaseServiceError, DatabaseCompatibilityError
from services.storage_diagnostics import database_failure
from test_meta_task_links import metadata
from test_meta_structural_validation import ios

OPERATION = '33333333-3333-4333-8333-333333333333'


class DatabaseError(RuntimeError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def client(ready=True):
    result = Mock()
    def rpc(name, params):
        data = {'api_version': 2, 'ready': ready, 'reasons': []} if name == 'lesson_plan_storage_status' else [{'lesson_id': 'stored', 'daily_count': 1}]
        return Mock(execute=Mock(return_value=Mock(data=data)))
    result.rpc.side_effect = rpc
    return result


class LessonStorageTest(unittest.TestCase):
    def test_missing_rpc_prevents_paid_generator_for_every_college(self):
        for college in ('META', 'Q', 'ETEC', 'ALT', 'ATKK'):
            with self.subTest(college=college):
                connection = client()
                connection.rpc.side_effect = DatabaseError('PGRST202', 'secret raw SQL')
                provider = Mock()
                with self.assertRaises(DatabaseCompatibilityError):
                    generate_lesson_plan(SupabaseService(connection), provider, college=college)
                provider.assert_not_called()
                connection.table.assert_not_called()
                self.assertEqual(connection.rpc.call_args.args, ('lesson_plan_storage_status', {}))

    def test_incompatible_contract_prevents_ai(self):
        for data in ({'api_version': 1, 'ready': True}, {'api_version': 2, 'ready': False}, {}, [], None):
            connection = Mock()
            connection.rpc.return_value.execute.return_value.data = data
            provider = Mock()
            with self.assertRaises(DatabaseCompatibilityError): generate_lesson_plan(SupabaseService(connection), provider)
            provider.assert_not_called()

    def test_network_or_permission_failure_never_calls_ai(self):
        for code in ('UNKNOWN', '42501', 'PGRST301', 'PGRST203'):
            connection = Mock()
            connection.rpc.side_effect = DatabaseError(code, 'private credentials')
            provider = Mock()
            with self.assertRaises(SupabaseServiceError): generate_lesson_plan(SupabaseService(connection), provider)
            provider.assert_not_called()

    def test_ready_storage_calls_provider_once(self):
        provider = Mock(return_value='ready')
        connection = client()
        result = generate_lesson_plan(SupabaseService(connection), provider, 'topic', college='Q', operation_id=OPERATION)
        self.assertEqual(result, 'ready')
        provider.assert_called_once_with('topic', college='Q')
        connection.table.assert_not_called()

    def test_local_diagnostics_do_not_contain_server_text(self):
        raw = 'sk-api-secret Bearer jwt-secret Teacher Иванов email@example.com SELECT * FROM profiles'
        with patch.dict(os.environ, {'LESSONAI_ENV': 'development'}, clear=True), self.assertLogs('services.storage_diagnostics') as logs:
            info = database_failure(DatabaseError('42702', raw), 'save_lesson_plan_with_daily_limit', 'save', OPERATION)
        self.assertEqual(info['code'], '42702')
        self.assertEqual(info['reason'], 'ambiguous_sql_reference')
        for secret in ('sk-api', 'jwt-secret', 'Иванов', 'email@example.com', 'SELECT'):
            self.assertNotIn(secret, str(logs.output))
        self.assertIn(OPERATION, str(logs.output))

    def test_production_does_not_log_detailed_failure(self):
        with patch.dict(os.environ, {'LESSONAI_ENV': 'production'}, clear=True), patch('services.storage_diagnostics.logging.getLogger') as logger:
            database_failure(DatabaseError('23514', 'private data'), 'save_lesson_plan_with_daily_limit', 'save', OPERATION)
        logger.assert_not_called()

    def test_error_codes_are_kept_but_public_message_is_safe(self):
        for code, message, reason in (('23514', 'private row', 'constraint_violation'),
                                      ('42702', 'private row', 'ambiguous_sql_reference'),
                                      ('P0001', 'IDEMPOTENCY_CONFLICT', 'idempotency_conflict')):
            connection = client()
            connection.rpc.side_effect = DatabaseError(code, message)
            with self.assertRaises(SupabaseServiceError) as error:
                SupabaseService(connection).save_lesson_plan(metadata(), '{}', operation_id=OPERATION)
            self.assertEqual(error.exception.diagnostics['code'], code)
            self.assertEqual(error.exception.diagnostics['reason'], reason)
            self.assertNotIn('private row', str(error.exception))
            connection.table.assert_not_called()

    def test_lost_response_retry_has_identical_key_and_payload_without_ai(self):
        connection = client()
        sent = []
        def rpc(name, params):
            if name == 'lesson_plan_storage_status':
                return Mock(execute=Mock(return_value=Mock(data={'api_version': 2, 'ready': True})))
            sent.append(copy.deepcopy(params))
            if len(sent) == 1: return Mock(execute=Mock(side_effect=TimeoutError('response lost')))
            return Mock(execute=Mock(return_value=Mock(data=[{'lesson_id': 'saved', 'daily_count': 1}])))
        connection.rpc.side_effect = rpc
        pending = PendingLessonPlan(ios(), metadata(), get_college_config('META')['template'], operation_id=OPERATION)
        service = SupabaseService(connection)
        with patch('services.ai.openai_generate') as paid:
            with self.assertRaises(SupabaseServiceError): pending.save(service)
            doc, saved = pending.save(service)
            paid.assert_not_called()
        self.assertEqual(sent[0], sent[1])
        self.assertEqual(sent[0]['p_operation_id'], OPERATION)
        self.assertEqual(saved['daily_count'], 1)
        self.assertEqual(len(Document(doc).tables[2].columns), 5)

    def test_saved_json_export_and_storage_for_all_colleges_without_ai(self):
        for college in ('META', 'Q', 'ETEC', 'ALT', 'ATKK'):
            with self.subTest(college=college):
                root = Path(__file__).resolve().parents[1]
                plan = ios() if college == 'META' else json.loads((root / 'docs/qa/q/lesson_q_ru.json').read_text())
                connection = client()
                service = SupabaseService(connection)
                pending = PendingLessonPlan(plan, metadata(college), get_college_config(college)['template'])
                with patch('services.ai.openai_generate') as paid:
                    document, _ = pending.save(service)
                    paid.assert_not_called()
                calls = [c for c in connection.rpc.call_args_list if c.args[0] == 'save_lesson_plan_with_daily_limit']
                self.assertEqual(len(calls), 1)
                payload = json.loads(calls[0].args[1]['p_material_content'])
                if college == 'META':
                    self.assertEqual(sum(i['points'] for i in payload['assessment_items']), 100)
                    self.assertEqual(sum(s['duration_minutes'] for s in payload['lesson_stages']), 70)
                if college == 'Q':
                    word = Document(document)
                    self.assertIn(payload['interdisciplinary_connections'], '\n'.join(c.text for table in word.tables for row in table.rows for c in row.cells))
                connection.table.assert_not_called()

    def test_pending_content_is_frozen_from_mutable_input(self):
        plan = ios()
        pending = PendingLessonPlan(plan, metadata(), get_college_config('META')['template'])
        plan['goal'] = 'Changed'
        self.assertNotEqual(json.loads(pending.content)['goal'], 'Changed')

    def test_schema_failure_after_ai_keeps_ready_content_for_unpaid_retry(self):
        connection = client()
        checks = iter([True, False, True])
        def rpc(name, params):
            data = {'api_version': 2, 'ready': next(checks)} if name == 'lesson_plan_storage_status' else [{'lesson_id': 'saved', 'daily_count': 1}]
            return Mock(execute=Mock(return_value=Mock(data=data)))
        connection.rpc.side_effect = rpc
        service = SupabaseService(connection)
        provider = Mock(return_value=json.dumps(ios()))
        result = generate_lesson_plan(service, provider, operation_id=OPERATION)
        pending = PendingLessonPlan(result, metadata(), get_college_config('META')['template'], operation_id=OPERATION)
        with self.assertRaises(DatabaseCompatibilityError): pending.save(service)
        self.assertEqual(json.loads(pending.content), ios())
        pending.save(service)
        provider.assert_called_once()
        writes = [call for call in connection.rpc.call_args_list if call.args[0] == 'save_lesson_plan_with_daily_limit']
        self.assertEqual(len(writes), 1)

    def test_free_form_diagnostic_identifiers_are_sanitized(self):
        info = database_failure(DatabaseError('IVANA', 'secret'), 'teacher@example.com', 'sk-key', 'Bearer token')
        self.assertEqual(info['rpc'], 'unknown_rpc')
        self.assertEqual(info['phase'], 'unknown_phase')
        self.assertEqual(info['code'], 'UNKNOWN')
        self.assertEqual(info['operation_id'], 'invalid_operation_id')

    def test_bad_uuid_is_rejected_before_rpc(self):
        connection = client()
        with self.assertRaises(SupabaseServiceError):
            SupabaseService(connection).save_lesson_plan(metadata(), '{}', operation_id='not-a-uuid')
        connection.rpc.assert_not_called()

