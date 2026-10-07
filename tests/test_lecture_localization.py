import unittest
from datetime import date
from unittest.mock import patch

from docx import Document

from services.docx_generator import build_lecture_docx
from services.lecture_localization import lecture_labels
from services.preview_renderer import _lecture_body, render_document_preview
from services.ai import rework_lecture


PAYLOAD = {
    "relevance": "Content",
    "learning_goals": ["Content"],
    "key_terms": [{"definition": "Content"}],
    "main_material": [{"content": "Content", "examples": ["Content"]}],
    "practice_examples": ["Content"],
    "case_study": dict.fromkeys(("scenario", "analysis", "solution"), "Content"),
    "common_errors": [dict.fromkeys(("error", "why", "fix"), "Content")],
    "summary": ["Content"],
    "review_questions": ["Content"],
    "preparation_for_practice": "Content",
}


class LectureLocalizationTest(unittest.TestCase):
    def test_export_and_preview_use_selected_language_for_every_label(self):
        body_keys = set(lecture_labels()) - {
            "lecture", "teacher", "date", "subject", "page", "schema", "visual", "logo"
        }
        for language in ("Русский", "Қазақша", "English"):
            with self.subTest(language=language):
                labels = lecture_labels(language)
                document = Document(build_lecture_docx(
                    PAYLOAD, language=language, teacher="Name", lesson_date=date(2026, 10, 7)
                ))
                text = "\n".join(p.text for p in document.paragraphs)
                footer = "\n".join(p.text for s in document.sections for p in s.footer.paragraphs)
                preview = _lecture_body(PAYLOAD, language)
                for key in body_keys:
                    self.assertIn(labels[key], text)
                    self.assertIn(labels[key], preview)
                self.assertIn(labels["lecture"].upper(), text)
                self.assertIn(labels["teacher"], text)
                self.assertIn(labels["date"], text)
                self.assertIn(labels["subject"], text)
                self.assertIn(labels["page"], footer)
                if language != "Русский":
                    for key, russian in lecture_labels().items():
                        if russian != labels[key]:
                            self.assertNotIn(russian, text + footer + preview)

    @patch("services.preview_renderer.st.markdown")
    def test_preview_cover_is_localized(self, markdown):
        render_document_preview(PAYLOAD, "Лекция", {"language": "English", "teacher": "Name"})
        html = markdown.call_args.args[0]
        self.assertIn("LECTURE", html)
        self.assertIn("Teacher:", html)
        self.assertIn("Date:", html)
        self.assertNotIn("Преподаватель", html)

    @patch("services.ai.openai_generate", return_value="{}")
    def test_rework_preserves_selected_language(self, generate):
        rework_lecture("{}", "Добавить примеры", None, language="Қазақша")
        self.assertIn("Язык лекции: Қазақша", generate.call_args.args[1])
