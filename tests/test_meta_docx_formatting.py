import re
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from config.colleges import get_college_config
from meta_fixtures import meta_plan
from services.docx_generator import build_doc, _fill_document, _document_paragraphs
from services.meta_formatting import dynamic_run_font
from services.meta_localization import META_LABELS
from services.preview_renderer import _meta_lesson_body

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = get_college_config('META')['template']


def create(language):
    return build_doc('План урока', meta_plan(language), template_path=TEMPLATE,
                     fields={'language': language, 'date': '2026-10-09', 'topic': 'Фишинг',
                             'subject': 'Информатика', 'teacher': 'Teacher', 'course': 1,
                             'group': 'IT-1', 'lesson_type': 'Дәріс' if language == 'Қазақша' else 'Лекция'})


def assert_font(test, run, size=None):
    fonts = run._r.rPr.rFonts
    for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        test.assertEqual(fonts.get(qn('w:' + script)), 'Times New Roman', run.text)
        test.assertIsNone(fonts.get(qn('w:' + script + 'Theme')))
    test.assertEqual(run.font.name, 'Times New Roman')
    test.assertIsNotNone(run.font.size)
    if size is not None:
        test.assertEqual(run.font.size.pt, size)


class MetaDocxFormattingTest(unittest.TestCase):
    def test_method_and_table_captions_follow_selected_language(self):
        for language, code in (('Русский', 'ru'), ('Қазақша', 'kk')):
            with self.subTest(language=language):
                doc = Document(create(language))
                methods = doc.tables[1].rows[9].cells[1].text
                labels = META_LABELS[code]
                self.assertIn(labels['explanatory_method_label'] + ':', methods)
                self.assertIn(labels['practical_task_label'] + ':', methods)
                if code == 'kk':
                    for key in ('explanatory_method_label', 'practical_task_label'):
                        self.assertNotIn(META_LABELS['ru'][key], methods)
                    self.assertNotIn('Этапы занятия', doc.tables[1].rows[13].cells[0].text)
                    self.assertNotIn('Запланированная деятельность', doc.tables[1].rows[13].cells[1].text)
                self.assertEqual(doc.tables[2].rows[0].cells[0].text, labels['headers'][0])
                self.assertEqual(doc.tables[2].rows[-1].cells[0].text, labels['total'])

    def test_all_populated_runs_use_times_new_roman_and_original_size(self):
        source = Document(TEMPLATE)
        for language in ('Русский', 'Қазақша'):
            with self.subTest(language=language):
                doc = Document(create(language))
                # Every placeholder paragraph, including method labels, is normalized.
                before = source.tables[1]
                after = doc.tables[1]
                for original_row, row in zip(before.rows[:14], after.rows[:14]):
                    for original_cell, cell in zip(original_row.cells, row.cells):
                        for original, paragraph in zip(original_cell.paragraphs, cell.paragraphs):
                            if re.search(r'\{\{?\w+\}?\}', original.text):
                                for run in paragraph.runs:
                                    if run.text:
                                        assert_font(self, run)
                for table, rows in ((doc.tables[1], doc.tables[1].rows[13:]),
                                    (doc.tables[2], doc.tables[2].rows)):
                    for row in rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    if run.text:
                                        assert_font(self, run, 12)
                self.assertTrue(doc.tables[2].rows[0].cells[0].paragraphs[0].runs[0].bold)
                self.assertTrue(doc.tables[1].rows[14].cells[0].paragraphs[0].runs[0].bold)

    def test_missing_font_uses_12pt_and_explicit_sizes_survive(self):
        doc = Document()
        doc.styles['Normal'].font.size = None
        run = doc.add_paragraph().add_run('Ә Ғ Қ Ң Ө Ұ Ү Һ І English')
        dynamic_run_font(run)
        assert_font(self, run, 12)
        run.font.size = Pt(14)
        run.bold = True
        fonts = run._r.rPr.rFonts
        for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + script + 'Theme'), 'minorHAnsi')
        dynamic_run_font(run)
        assert_font(self, run, 14)
        self.assertTrue(run.bold)

    def test_inherited_explicit_template_size_is_preserved(self):
        doc = Document()
        doc.styles['Normal'].font.size = Pt(13)
        run = doc.add_paragraph().add_run('Text')
        dynamic_run_font(run)
        assert_font(self, run, 13)

    def test_rubric_widths_match_grid_cells_and_fit_page(self):
        doc = Document(create('Русский'))
        table = doc.tables[2]
        source = Document(ROOT / 'docs/template_backups/template_meta.before_localization.docx')
        self.assertEqual(len(table.columns), 5)
        widths = [column.width for column in table.columns]
        self.assertGreater(widths[3], source.tables[2].columns[3].width * 2)
        available = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
        self.assertLessEqual(sum(widths), available)
        self.assertLess(available - sum(widths), 635)
        self.assertFalse(table.autofit)
        for row in table.rows:
            self.assertEqual([cell.width for cell in row.cells], widths)
            self.assertTrue(row._tr.xpath('./w:trPr/w:cantSplit'))
            self.assertFalse(row._tr.xpath('./w:trPr/w:trHeight'))
        self.assertEqual(table.rows[1].cells[4].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(table.rows[1].cells[2].paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_descriptors_are_separate_hanging_paragraphs_without_loss(self):
        payload = meta_plan('Қазақша')
        doc = Document(create('Қазақша'))
        for item, row in zip(payload['assessment_items'], doc.tables[2].rows[1:-1]):
            paragraphs = row.cells[3].paragraphs
            self.assertEqual(len(paragraphs), len(item['descriptors']))
            for descriptor, paragraph in zip(item['descriptors'], paragraphs):
                self.assertIn(descriptor['text'], paragraph.text)
                self.assertTrue(paragraph.text.startswith('• '))
                self.assertLess(paragraph.paragraph_format.first_line_indent, 0)
                self.assertTrue(paragraph.text.strip())
        self.assertEqual(sum(int(r.cells[4].text) for r in doc.tables[2].rows[1:-1]), 100)
        self.assertEqual(doc.tables[1].cell(0, 3).text, '09.10.2026 ж./г.')
        minutes = [int(re.search(r'\((\d+) мин', r.cells[0].text).group(1)) for r in doc.tables[1].rows[14:]]
        self.assertEqual(sum(minutes), 70)
        self.assertIn('ә', doc.tables[1].rows[9].cells[1].text)

    def test_fixed_shell_logos_sections_and_package_parts_are_preserved(self):
        before_path = ROOT / 'docs/template_backups/template_meta.before_localization.docx'
        before = Document(before_path)
        after = Document(create('Қазақша'))
        self.assertEqual([p._p.xml for p in before.paragraphs], [p._p.xml for p in after.paragraphs])
        self.assertEqual(before.tables[0]._tbl.xml, after.tables[0]._tbl.xml)
        self.assertEqual([s._sectPr.xml for s in before.sections], [s._sectPr.xml for s in after.sections])
        self.assertEqual([c.width for c in before.tables[1].columns], [c.width for c in after.tables[1].columns])
        with ZipFile(before_path) as reference, ZipFile(create('Қазақша')) as output:
            for name in reference.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(reference.read(name), output.read(name), name)

    def test_no_global_translation_of_user_content(self):
        payload = meta_plan()
        literal = 'Цитата: Практическое задание: не менять.'
        payload['resources'] = literal
        payload['explanatory_illustrative_method'] = literal
        doc = Document(build_doc('План урока', payload, template_path=TEMPLATE,
                                 fields={'language': 'Русский', 'date': '2026-10-09'}))
        self.assertIn(literal, doc.tables[1].rows[9].cells[1].text)
        self.assertIn(literal, doc.tables[1].rows[11].cells[1].text)

    def test_other_college_placeholders_keep_original_formatting(self):
        for college in ('ETEC', 'ALT', 'Q', 'ATKK'):
            with self.subTest(college=college):
                doc = Document(get_college_config(college)['template'])
                originals = list(_document_paragraphs(doc))
                signatures = [[(run.font.name, run.font.size, run.bold, run.italic) for run in p.runs]
                              for p in originals]
                _fill_document(doc, {'topic': 'Тест', 'subject': 'Тест', 'date': '09.10.2026'})
                self.assertEqual(signatures, [[(run.font.name, run.font.size, run.bold, run.italic) for run in p.runs]
                                               for p in originals])

    def test_preview_uses_centralized_method_labels(self):
        html = _meta_lesson_body(meta_plan('Қазақша'), {'language': 'Қазақша', 'date': '2026-10-09'})
        self.assertIn('Түсіндіру-көрнекілік әдісі:', html)
        self.assertIn('Практикалық тапсырма:', html)
        self.assertNotIn('Практическое задание:', html)
