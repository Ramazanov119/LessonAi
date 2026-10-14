import copy
import json
import unittest
from datetime import date
from unittest.mock import Mock, patch

from docx import Document
from config.colleges import get_college_config
from models import LessonMetadata
from meta_fixtures import meta_plan
from services.ai import create_lesson, AIServiceError
from services.docx_generator import build_doc
from services.lesson_validation import normalize_meta_lesson, LessonValidationError
from services.lesson_plan_service import save_generated_lesson_plan
from services.supabase_service import SupabaseService, SupabaseServiceError

TASK = 'Ответить на вопросы по структуре мобильного приложения и жизненному циклу Activity/Screen'
ACTIVITY = 'Преподаватель проводит опрос: студенты отвечают на вопросы о структуре мобильного приложения и жизненном цикле Activity / Screen.'


def navigation_plan(with_ids=False):
    payload = meta_plan()
    payload['goal'] = 'Научить реализовывать навигацию страниц мобильного приложения'
    actualization = next(s for s in payload['lesson_stages'] if s['field_key'] == 'actualization')
    actualization['content'] = ACTIVITY
    payload['actualization'] = ACTIVITY
    item = payload['assessment_items'][0]
    item.update(task_name=TASK, stage_key='actualization', bloom_level='Знание',
                assessment_criterion='Описывает структуру мобильного приложения и жизненный цикл Activity/Screen',
                descriptors=[{'text': 'Называет компоненты структуры мобильного приложения', 'points': 10},
                             {'text': 'Объясняет жизненный цикл Activity/Screen', 'points': 10}])
    actualization['tasks'] = [{'id': 'mobile_recall', 'name': 'Опрос по мобильному приложению' if with_ids else TASK, 'instruction': ACTIVITY}]
    if with_ids:
        item['task_id'] = 'mobile_recall'
    else:
        item.pop('task_id', None)
    return payload


def metadata(college='META'):
    return LessonMetadata(full_name='Teacher', college=college, subject='Мобильная разработка',
                          topic='Навигация страниц', course=1, group_name='ППО', language='Русский',
                          lesson_type='Комбинированный урок', duration=70, lesson_date=date(2026, 10, 13),
                          speciality='IT', chair='Chair')


class MetaTaskLinksTest(unittest.TestCase):
    def test_reported_navigation_question_is_recovered_from_existing_activity(self):
        plan = navigation_plan()
        self.assertNotIn(TASK.casefold(), ACTIVITY.casefold())  # Old predicate rejected this valid activity.
        repaired = normalize_meta_lesson(plan, lesson_type='Комбинированный урок', lesson_date=date(2026, 10, 13))
        item = repaired['assessment_items'][0]
        self.assertEqual(item['stage_key'], 'actualization')
        self.assertEqual(item['task_id'], 'mobile_recall')
        stage = next(s for s in repaired['lesson_stages'] if s['field_key'] == item['stage_key'])
        self.assertEqual(stage['tasks'][0]['instruction'], ACTIVITY)
        self.assertEqual(sum(s['duration_minutes'] for s in repaired['lesson_stages']), 70)
        self.assertEqual(sum(i['points'] for i in repaired['assessment_items']), 100)
        self.assertEqual(repaired, normalize_meta_lesson(repaired, lesson_date=date(2026, 10, 13)))
        doc = Document(build_doc('План урока', repaired, template_path=get_college_config('META')['template'],
                                 fields={'language': 'Русский', 'date': '2026-10-13'}))
        self.assertIn(TASK, doc.tables[2].rows[1].cells[0].text)
        self.assertIn(ACTIVITY, doc.tables[1].rows[15].cells[1].text)

    def test_ids_link_different_titles_without_inventing_rubric_activity(self):
        plan = navigation_plan(True)
        self.assertNotIn('Опрос по мобильному приложению', ACTIVITY)
        result = normalize_meta_lesson(plan)
        item = result['assessment_items'][0]
        self.assertEqual(item['task_id'], 'mobile_recall')
        self.assertEqual(item['task_name'], 'Опрос по мобильному приложению')
        self.assertEqual(result['actualization'], ACTIVITY)

    def test_declared_instruction_is_rendered_when_not_in_stage_prose(self):
        plan = navigation_plan(True)
        plan['lesson_stages'][1]['content'] = 'Введение в навигацию страниц мобильного приложения.'
        result = normalize_meta_lesson(plan)
        self.assertIn(ACTIVITY, result['actualization'])
        self.assertEqual(result['assessment_items'][0]['task_id'], 'mobile_recall')

    def test_missing_name_is_repaired_but_wrong_stage_is_rejected(self):
        plan = navigation_plan(True)
        del plan['assessment_items'][0]['task_name']
        result = normalize_meta_lesson(plan)
        self.assertEqual(result['assessment_items'][0]['task_name'], 'Опрос по мобильному приложению')
        plan['assessment_items'][0]['stage_key'] = 'practice'
        with self.assertRaises(LessonValidationError):
            normalize_meta_lesson(plan)

    def test_dict_and_alias_variants_are_repaired_once(self):
        for shape in ('single', 'mapping'):
            with self.subTest(shape=shape):
                plan = navigation_plan(True)
                stage = plan['lesson_stages'][1]
                task = {'task_id': 'mobile_recall', 'task_name': 'Опрос по мобильному приложению', 'instruction': ACTIVITY}
                stage['tasks'] = task if shape == 'single' else {'mobile_recall': task}
                result = normalize_meta_lesson(plan)
                self.assertEqual(result['assessment_items'][0]['task_id'], result['lesson_stages'][1]['tasks'][0]['id'])

    def test_unknown_id_absent_task_and_wrong_stage_are_rejected(self):
        for variant in ('unknown_id', 'absent_task', 'wrong_stage'):
            with self.subTest(variant=variant):
                plan = navigation_plan(True)
                if variant == 'unknown_id':
                    plan['assessment_items'][0]['task_id'] = 'not_in_lesson'
                elif variant == 'absent_task':
                    plan['lesson_stages'][1]['tasks'] = []
                else:
                    plan['assessment_items'][0]['stage_key'] = 'practice'
                with self.assertRaises(LessonValidationError):
                    normalize_meta_lesson(plan)

    def test_ambiguous_missing_id_is_not_guessed(self):
        plan = navigation_plan()
        plan['lesson_stages'][1]['tasks'].append({'id': 'another', 'name': 'Другой опрос', 'instruction': 'Отвечает на дополнительные вопросы'})
        with self.assertRaises(LessonValidationError):
            normalize_meta_lesson(plan)

    def test_duplicate_assessment_id_and_duplicate_declared_ids_are_rejected(self):
        plan = navigation_plan(True)
        plan['assessment_items'].append(copy.deepcopy(plan['assessment_items'][0]))
        with self.assertRaisesRegex(LessonValidationError, 'оценено повторно'):
            normalize_meta_lesson(plan)
        plan = navigation_plan(True)
        plan['lesson_stages'][2]['tasks'] = copy.deepcopy(plan['lesson_stages'][1]['tasks'])
        with self.assertRaisesRegex(LessonValidationError, 'Идентификатор задания повторяется'):
            normalize_meta_lesson(plan)

    @patch('services.ai.openai_generate')
    def test_reported_manual_generation_needs_one_paid_provider_call_only(self, provider):
        provider.return_value = json.dumps(navigation_plan())
        result = json.loads(create_lesson('Мобильная разработка', 'Навигация страниц', 'Русский', 'IT', 'Chair',
                                         'Комбинированный урок', 70, None, college='META', lesson_date=date(2026, 10, 13)))
        provider.assert_called_once()
        self.assertIn('task_id', provider.call_args.args[0])
        self.assertEqual(result['assessment_items'][0]['stage_key'], 'actualization')
        self.assertEqual(result['lesson_duration'], 70)
        self.assertEqual(sum(item['points'] for item in result['assessment_items']), 100)

    @patch('services.ai.openai_generate')
    def test_other_colleges_do_not_run_meta_validation_even_with_bad_assessment(self, provider):
        for college in ('ETEC', 'ALT', 'Q', 'ATKK'):
            with self.subTest(college=college), patch('services.lesson_validation.normalize_meta_lesson') as validator:
                provider.return_value = json.dumps({'goal': 'Научить навигации', 'assessment_items': [
                    {'task_name': TASK, 'stage_key': 'nonexistent'}],
                    'interdisciplinary_connections': 'Логика — анализ условий перехода между страницами'})
                result = create_lesson('Мобильная разработка', 'Навигация страниц', 'Русский', 'IT', 'Chair',
                                       'Комбинированный урок', 90, None, college=college)
                validator.assert_not_called()
                with patch('services.docx_generator.normalize_meta_lesson') as exporter_validator:
                    Document(build_doc('План урока', result, template_path=get_college_config(college)['template'],
                                       fields={'language': 'Русский', 'date': '2026-10-13'}, college=college))
                    exporter_validator.assert_not_called()

    def test_non_meta_save_pipeline_does_not_apply_meta_assessment_rules(self):
        for college in ('ETEC', 'ALT', 'Q', 'ATKK'):
            with self.subTest(college=college), patch('services.lesson_plan_service.normalize_meta_lesson') as validator:
                service = Mock()
                service.save_lesson_plan.return_value = {'lesson_id': 'new', 'daily_count': 1}
                payload = {'goal': 'Навигация страниц', 'assessment_items': [{'task_name': TASK}],
                           'interdisciplinary_connections': 'Логика — анализ условий перехода между страницами'}
                document, _ = save_generated_lesson_plan(service, payload, metadata(college), get_college_config(college)['template'])
                validator.assert_not_called()
                self.assertEqual(len(Document(document).tables), len(Document(get_college_config(college)['template']).tables))
                service.save_lesson_plan.assert_called_once()

    def test_invalid_export_neither_saves_material_nor_consumes_quota(self):
        service = Mock()
        bad = navigation_plan(True)
        bad['assessment_items'][0]['task_id'] = 'missing'
        with self.assertRaises(LessonValidationError):
            save_generated_lesson_plan(service, bad, metadata(), get_college_config('META')['template'])
        service.save_lesson_plan.assert_not_called()
        service.create_lesson.assert_not_called()
        service.upsert_material.assert_not_called()

    @patch('services.lesson_plan_service.build_doc', side_effect=RuntimeError('Word export failed'))
    def test_docx_failure_prevents_any_database_write(self, export):
        service = Mock()
        with self.assertRaises(RuntimeError):
            save_generated_lesson_plan(service, navigation_plan(), metadata(), get_college_config('META')['template'])
        service.save_lesson_plan.assert_not_called()

    def test_valid_plan_uses_one_atomic_rpc_and_existing_lesson_is_passed(self):
        client = Mock()
        client.rpc.side_effect = lambda name, params: Mock(execute=Mock(return_value=Mock(data=({'api_version': 2, 'ready': True, 'reasons': []} if name == 'lesson_plan_storage_status' else [{'lesson_id': 'existing', 'daily_count': 3}]))))
        doc, saved = save_generated_lesson_plan(SupabaseService(client), navigation_plan(), metadata(),
                                               get_college_config('META')['template'], existing_lesson_id='11111111-1111-4111-8111-111111111111')
        self.assertEqual(saved['daily_count'], 3)
        self.assertEqual(len(Document(doc).tables), 3)
        self.assertEqual(client.rpc.call_count, 2)
        name, params = client.rpc.call_args.args
        self.assertEqual(name, 'save_lesson_plan_with_daily_limit')
        self.assertEqual(params['p_existing_lesson_id'], '11111111-1111-4111-8111-111111111111')
        self.assertEqual(params['p_topic'], 'Навигация страниц')
        self.assertEqual(params['p_group_name'], 'ППО')
        client.table.assert_not_called()

    def test_saved_json_matches_corrected_task_ids_timing_and_rubric(self):
        service = Mock()
        service.save_lesson_plan.return_value = {'lesson_id': 'new', 'daily_count': 1}
        plan = navigation_plan()
        for stage in plan['lesson_stages']:
            stage['duration_minutes'] = 5
        document, _ = save_generated_lesson_plan(service, plan, metadata(), get_college_config('META')['template'])
        saved = json.loads(service.save_lesson_plan.call_args.args[1])
        self.assertEqual(sum(stage['duration_minutes'] for stage in saved['lesson_stages']), 70)
        self.assertEqual(sum(item['points'] for item in saved['assessment_items']), 100)
        self.assertTrue(all(item['task_id'] for item in saved['assessment_items']))
        self.assertEqual(saved['date'], '2026-10-13')
        table = Document(document).tables[2]
        self.assertEqual([int(row.cells[4].text) for row in table.rows[1:-1]],
                         [item['points'] for item in saved['assessment_items']])

    def test_save_failure_does_not_fall_back_to_non_atomic_creation(self):
        client = Mock()
        client.rpc.return_value.execute.side_effect = RuntimeError('PGRST202: function missing')
        with self.assertRaisesRegex(SupabaseServiceError, 'Требуется обновление базы'):
            save_generated_lesson_plan(SupabaseService(client), navigation_plan(), metadata(), get_college_config('META')['template'])
        client.rpc.assert_called_once()
        client.table.assert_not_called()
