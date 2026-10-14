import ast
import copy
import json
import re
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock, patch
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Pt

from config.colleges import COLLEGES, get_college_config
from services.ai import create_lesson, AIServiceError
from services.docx_generator import build_doc, _fill_document
from services.lesson_validation import (
    LessonValidationError, normalize_meta_lesson, lesson_date_value, almaty_today,
)
from services.meta_docx import find_meta_tables
from services.preview_renderer import _meta_lesson_body
from services.supabase_service import SupabaseService
from models import LessonMetadata
from meta_fixtures import meta_plan

ROOT = Path(__file__).resolve().parents[1]


def export(payload=None, **fields):
    return Document(build_doc('План урока', payload or meta_plan(),
                             template_path=get_college_config('META')['template'],
                             fields={'date': date(2026, 10, 9), 'language': 'Русский', **fields}))


class MetaLessonTest(unittest.TestCase):
    def test_russian_and_kazakh_assessment_is_filled_without_duplicate_table(self):
        for language, header in (("Русский", "Задание"), ("Қазақша", "Тапсырма")):
            with self.subTest(language=language):
                payload = meta_plan(language)
                doc = export(payload, language=language)
                self.assertEqual(len(doc.tables), 3)
                lesson, rubric = find_meta_tables(doc)
                self.assertEqual(len(rubric.rows), len(payload['assessment_items']) + 2)
                self.assertEqual(rubric.cell(0, 0).text, header)
                self.assertEqual(sum(int(row.cells[-1].text) for row in rubric.rows[1:-1]), 100)
                for row, item in zip(rubric.rows[1:-1], payload['assessment_items']):
                    self.assertEqual(row.cells[0].text, item['task_name'])
                    self.assertTrue(row.cells[2].text.strip())
                    self.assertTrue(row.cells[3].text.strip())
                    self.assertIn(item['task_name'], '\n'.join(c.text for r in lesson.rows for c in r.cells))
                if language == 'Қазақша':
                    self.assertNotIn('Знание', ' '.join(c.text for r in rubric.rows for c in r.cells))

    def test_positive_weights_are_normalized_without_changing_source(self):
        payload = meta_plan()
        for item in payload['assessment_items']:
            item['points'] = 17
        result = normalize_meta_lesson(payload)
        self.assertEqual(sum(item['points'] for item in result['assessment_items']), 100)
        self.assertEqual(payload['assessment_items'][0]['points'], 17)

    def test_unrelated_task_is_rejected(self):
        payload = meta_plan()
        payload['assessment_items'][0]['task_id'] = 'task_not_declared'
        with self.assertRaises(LessonValidationError):
            export(payload)

    def test_criterion_for_other_task_is_rejected(self):
        payload = meta_plan()
        payload['assessment_items'][0]['assessment_criterion'] = {'task_id': 'task_2', 'text': 'Обосновывает выбор'}
        with self.assertRaises(LessonValidationError):
            export(payload)

    def test_empty_criteria_descriptors_and_invalid_scores_are_rejected(self):
        for key, value in (('assessment_criterion', ''), ('descriptors', []), ('points', -1), ('points', '1.5'), ('points', True), ('bloom_level', 'Синтез')):
            with self.subTest(key=key, value=value):
                payload = meta_plan()
                payload['assessment_items'][0][key] = value
                with self.assertRaises(LessonValidationError):
                    export(payload)

    def test_missing_assessment_fails_instead_of_exporting_sample(self):
        payload = meta_plan()
        del payload['assessment_items']
        with self.assertRaisesRegex(LessonValidationError, 'Создайте план заново'):
            export(payload)

    def test_selected_date_overrides_all_ai_aliases(self):
        doc = export()
        text = doc.tables[1].cell(0, 3).text
        self.assertEqual(text, '09.10.2026 ж./г.')
        self.assertNotIn('28.06.2024', text)
        self.assertEqual(export(date='09.10.2026 ж./г.').tables[1].cell(0, 3).text, text)

    @patch('services.lesson_validation.almaty_today', return_value=date(2026, 10, 9))
    def test_missing_date_uses_current_almaty_date(self, today):
        self.assertEqual(export(date=None).tables[1].cell(0, 3).text, '09.10.2026 ж./г.')
        today.assert_called()

    def test_almaty_day_at_utc_boundary(self):
        with patch('services.lesson_validation.datetime') as clock:
            clock.now.return_value = datetime(2026, 10, 10, 0, 30, tzinfo=ZoneInfo('Asia/Almaty'))
            self.assertEqual(almaty_today(), date(2026, 10, 10))
            clock.now.assert_called_once_with(ZoneInfo('Asia/Almaty'))

    def test_saved_lesson_download_preserves_date_without_importing_app(self):
        tree = ast.parse((ROOT / 'app.py').read_text())
        function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == '_material_download')
        namespace = {'build_doc': build_doc, 'get_college_config': get_college_config}
        exec(compile(ast.Module(body=[function], type_ignores=[]), str(ROOT / 'app.py'), 'exec'), namespace)
        lesson = dict(topic='Фишинг', subject='Информатика', full_name='Teacher', lesson_date='2026-10-09',
                      course=1, group_name='IT-1', lesson_type='Лекция', language='Русский', college='META')
        _, file, _ = namespace['_material_download'](lesson, {'material_type': 'lesson_plan', 'content': json.dumps(meta_plan())})
        self.assertEqual(Document(file).tables[1].cell(0, 3).text, '09.10.2026 ж./г.')

    def test_lecture_practice_and_laboratory_timings_match_word(self):
        for kind in ('Лекция', 'Практическое занятие', 'Лабораторная работа'):
            with self.subTest(kind=kind):
                payload = meta_plan(lesson_type=kind)
                # Deliberately recover the old 60-minute distribution.
                for stage, value in zip(payload['lesson_stages'], (5, 5, 20, 20, 5, 2, 3)):
                    stage['duration_minutes'] = value
                result = normalize_meta_lesson(payload, lesson_type=kind, lesson_date=date(2026, 10, 9))
                self.assertEqual(result['lesson_duration'], 70)
                self.assertEqual(sum(s['duration_minutes'] for s in result['lesson_stages']), 70)
                self.assertTrue(all(s['duration_minutes'] > 0 for s in result['lesson_stages']))
                doc = export(result, lesson_type=kind)
                rows = doc.tables[1].rows[14:]
                word_minutes = [int(re.search(r'\((\d+) мин', row.cells[0].text).group(1)) for row in rows]
                self.assertEqual(word_minutes, [s['duration_minutes'] for s in result['lesson_stages']])
                self.assertNotEqual(word_minutes, [5, 5, 20, 20, 5, 2, 3])

    def test_variable_stage_and_assessment_counts(self):
        payload = meta_plan()
        payload['lesson_stages'] = [s for s in payload['lesson_stages'] if s['field_key'] not in ('org_steps', 'homework')]
        payload['assessment_items'] = payload['assessment_items'][:3]
        for item, points in zip(payload['assessment_items'], (40, 30, 30)):
            item['points'] = points
            for descriptor in item['descriptors']:
                descriptor['points'] = points // 2
        doc = export(payload)
        self.assertEqual(len(doc.tables[1].rows), 14 + 5)
        self.assertEqual(len(doc.tables[2].rows), 1 + 3 + 1)

    def test_old_plan_without_declared_tasks_does_not_fabricate_tasks(self):
        legacy = meta_plan()
        del legacy['lesson_stages']
        with self.assertRaises(LessonValidationError):
            normalize_meta_lesson(legacy, lesson_type='Практическое занятие')

    def test_negative_timing_rejected_and_zero_corrected(self):
        payload = meta_plan()
        payload['lesson_stages'][0]['duration_minutes'] = -1
        with self.assertRaises(LessonValidationError):
            export(payload)
        payload['lesson_stages'][0]['duration_minutes'] = 0
        result = normalize_meta_lesson(payload)
        self.assertGreater(result['lesson_stages'][0]['duration_minutes'], 0)
        self.assertEqual(sum(s['duration_minutes'] for s in result['lesson_stages']), 70)

    @patch('services.ai.openai_generate')
    def test_one_ai_call_includes_meta_schema_and_70_minutes(self, provider):
        provider.return_value = json.dumps(meta_plan())
        result = json.loads(create_lesson('Информатика', 'Фишинг', 'Русский', 'IT', 'Chair',
                                         'Лекция', 60, None, college='META', lesson_date=date(2026, 10, 9)))
        provider.assert_called_once()
        prompt = provider.call_args.args[0]
        self.assertIn('assessment_items', prompt)
        self.assertIn('lesson_stages', prompt)
        self.assertIn('70 минут', prompt)
        self.assertNotIn('60 минут', prompt)
        self.assertNotIn('Создай учебно-методическую лекцию', prompt)
        self.assertEqual(result['date'], '2026-10-09')
        self.assertEqual(result['lesson_duration'], 70)

    @patch('services.ai.openai_generate', return_value='{}')
    def test_invalid_ai_plan_is_not_accepted(self, provider):
        with self.assertRaises(AIServiceError):
            create_lesson('IT', 'Topic', 'Русский', 'IT', 'Chair', 'Лекция', 70, None, college='META')

    def test_saved_metadata_contains_selected_date_and_meta_70_minutes(self):
        client = Mock()
        client.rpc.return_value.execute.return_value.data = [{'lesson_id': 'fake'}]
        lesson = LessonMetadata(full_name='Teacher', college='META', subject='IT', topic='Фишинг',
                                group_name='IT-1', course=1, duration=70, lesson_date=date(2026, 10, 9),
                                language='Русский', lesson_type='Лекция', speciality='IT', chair='Chair')
        SupabaseService(client).create_lesson(lesson)
        parameters = client.rpc.call_args.args[1]
        self.assertEqual(parameters['p_lesson_date'], '2026-10-09')
        self.assertEqual(parameters['p_duration'], '70 минут')

    def test_validated_json_can_be_saved_without_database_migration(self):
        payload = normalize_meta_lesson(meta_plan(), lesson_date=date(2026, 10, 9))
        client = Mock()
        client.table.return_value.upsert.return_value.execute.return_value.data = [{'id': 'fake'}]
        SupabaseService(client).upsert_material('fake-lesson', 'lesson_plan', json.dumps(payload, ensure_ascii=False))
        stored = client.table.return_value.upsert.call_args.args[0]
        decoded = json.loads(stored['content'])
        self.assertEqual(decoded['lesson_stages'], payload['lesson_stages'])
        self.assertEqual(decoded['assessment_items'], payload['assessment_items'])

    def test_preview_uses_the_same_corrected_values(self):
        payload = meta_plan()
        payload['lesson_stages'][0]['duration_minutes'] = 1
        result = normalize_meta_lesson(payload, lesson_date=date(2026, 10, 9))
        html = _meta_lesson_body(payload, {'language': 'Русский', 'date': '2026-10-09'})
        for stage in result['lesson_stages']:
            self.assertIn(f"({stage['duration_minutes']} мин.)", html)
        self.assertIn('70 мин.', html)
        self.assertIn('<td>100</td>', html)

    def test_template_geometry_borders_and_other_colleges_preserved(self):
        original = Document(ROOT / 'docs/template_backups/template_meta.original.docx')
        result = export()
        for before, after in zip(original.tables[:2], result.tables[:2]):
            self.assertEqual([col.width for col in before.columns], [col.width for col in after.columns])
        self.assertEqual([x.xml for x in original.tables[2].rows[1].cells[0]._tc.xpath('./w:tcPr/w:tcBorders')],
                         [x.xml for x in result.tables[2].rows[1].cells[0]._tc.xpath('./w:tcPr/w:tcBorders')])
        for code, expected in (('META', 70), ('ETEC', 70), ('ALT', 90), ('Q', 90), ('ATKK', 70)):
            self.assertEqual(COLLEGES[code]['lesson_duration'], expected)
        for code in ('ETEC', 'ALT', 'Q', 'ATKK'):
            doc = Document(build_doc('План урока', {'date': '28.06.2024', 'interdisciplinary_connections': 'Математика — вычисления по теме'},
                                    template_path=COLLEGES[code]['template'], fields={'date': '2026-10-09'}))
            self.assertEqual(len(doc.tables), len(Document(COLLEGES[code]['template']).tables))

    def test_untouched_template_parts_remain_byte_identical(self):
        original = ROOT / 'docs/template_backups/template_meta.original.docx'
        template = ROOT / 'template_meta.docx'
        output = build_doc('План урока', meta_plan(), template_path=template,
                           fields={'date': '2026-10-09', 'language': 'Русский'})
        with ZipFile(original) as source, ZipFile(template) as edited, ZipFile(output) as generated:
            for name in source.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(source.read(name), edited.read(name), name)
                    self.assertEqual(edited.read(name), generated.read(name), name)

    def test_long_descriptors_wrap_without_fixed_row_height(self):
        payload = meta_plan()
        long_text = 'Проверяет адрес отправителя фишингового письма. ' * 20
        payload['assessment_items'][0]['descriptors'][0]['text'] = long_text
        doc = export(payload)
        row = doc.tables[2].rows[1]
        self.assertIn(long_text.strip(), row.cells[3].text)
        self.assertFalse(row._tr.xpath('./w:trPr/w:trHeight'))
        self.assertFalse(row.cells[3]._tc.xpath('./w:tcPr/w:noWrap'))
        available = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
        self.assertLessEqual(sum(column.width for column in doc.tables[2].columns), available)

    def test_split_run_placeholders_and_regulatory_dates(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('{{lesson_').bold = True
        p.add_run('date}} / {da')
        p.add_run('ta} Приказ 28.06.2024')
        p.runs[0].font.size = Pt(14)
        doc.sections[0].header.paragraphs[0].text = 'Дата: {date}'
        _fill_document(doc, {'lesson_date': '09.10.2026 ж./г.', 'data': '09.10.2026', 'date': '09.10.2026'})
        self.assertEqual(p.text, '09.10.2026 ж./г. / 09.10.2026 Приказ 28.06.2024')
        self.assertTrue(p.runs[0].bold)
        self.assertEqual(p.runs[0].font.size, Pt(14))
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, 'Дата: 09.10.2026')


if __name__ == '__main__':
    unittest.main()
