import copy
import json
import re
import unittest
from datetime import date
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from config.colleges import get_college_config
from services.ai import create_lesson
from services.docx_generator import build_doc
from services.lesson_plan_service import save_generated_lesson_plan
from services.lesson_validation import normalize_meta_lesson, validate_meta_lesson, LessonValidationError
from services.meta_diagnostics import anonymized_snapshot
from services.meta_scores import allocate_scores
from services.preview_renderer import _meta_lesson_body
from services.supabase_service import SupabaseService
from test_meta_structural_validation import ios
from test_meta_task_links import metadata


def weighted(values, code='ru'):
    plan = ios(code)
    plan['assessment_items'] = plan['assessment_items'][:len(values)]
    for item, value in zip(plan['assessment_items'], values):
        item['points'] = value
    return plan


class MetaScoresTest(unittest.TestCase):
    def assert_scores(self, weights, expected):
        source = weighted(weights)
        before = copy.deepcopy(source)
        result = normalize_meta_lesson(source)
        self.assertEqual([i['points'] for i in result['assessment_items']], expected)
        self.assertEqual(source, before)
        self.assertEqual(result, normalize_meta_lesson(result))
        self.assertEqual(sum(i['points'] for i in result['assessment_items']), 100)
        for item in result['assessment_items']:
            self.assertEqual(sum(item['descriptor_points']), item['points'])
        return result

    def test_ai_total_80(self): self.assert_scores([16] * 5, [20] * 5)
    def test_ai_total_95(self): self.assert_scores([20, 20, 20, 20, 15], [21, 21, 21, 21, 16])
    def test_ai_total_105(self): self.assert_scores([20, 20, 20, 20, 25], [19, 19, 19, 19, 24])
    def test_ai_total_100(self): self.assert_scores([10, 20, 30, 40], [10, 20, 30, 40])
    def test_requested_90_example(self): self.assert_scores([20, 20, 20, 20, 10], [23, 22, 22, 22, 11])
    def test_requested_equal_weights(self): self.assert_scores([15] * 4, [25] * 4)
    def test_fractional_weights(self): self.assert_scores([.5, 1.5, 2, 3, 3], [5, 15, 20, 30, 30])

    def test_missing_weights_use_explicit_bloom_configuration(self):
        for code, language in (('ru', 'Русский'), ('kk', 'Қазақша')):
            plan = ios(code)
            for item in plan['assessment_items']:
                item.pop('points')
                for descriptor in item['descriptors']: descriptor.pop('points')
            result = normalize_meta_lesson(plan, language=language)
            self.assertEqual([i['points'] for i in result['assessment_items']], [6, 24, 18, 23, 29])
            self.assertEqual(sum(i['points'] for i in result['assessment_items']), 100)

    def test_invalid_weights_are_not_repaired(self):
        for value in (-1, 0, True, '20', None, float('nan'), float('inf'), 10**1000):
            for position in ('task', 'descriptor'):
                with self.subTest(value_type=type(value).__name__, position=position):
                    plan = ios()
                    target = plan['assessment_items'][0]
                    if position == 'descriptor': target = target['descriptors'][0]
                    target['points'] = value
                    with self.assertRaises(LessonValidationError) as error: normalize_meta_lesson(plan)
                    self.assertEqual(error.exception.code, 'invalid_score_weight')
                    self.assertEqual(error.exception.task_id, 'task_1')

    def test_partially_missing_weights_use_configuration_in_matching_units(self):
        plan = ios()
        plan['assessment_items'][0]['points'] = 7.5
        for item in plan['assessment_items'][1:]: item.pop('points')
        plan['assessment_items'][0]['descriptors'][1].pop('points')
        result = normalize_meta_lesson(plan)
        self.assertEqual([i['points'] for i in result['assessment_items']], [6, 24, 18, 23, 29])
        self.assertEqual(result['assessment_items'][0]['descriptor_points'], [3, 3])

    def test_inconsistent_criteria_alias_cannot_be_saved(self):
        plan = normalize_meta_lesson(ios())
        plan['evaluation_criteria'] = 'Независимый критерий, отсутствующий в таблице.'
        service = Mock()
        with self.assertRaises(LessonValidationError) as error:
            save_generated_lesson_plan(service, plan, metadata(), get_college_config('META')['template'])
        self.assertEqual(error.exception.code, 'criteria_structure_mismatch')
        service.save_lesson_plan.assert_not_called()

    def test_bounds_and_impossible_allocation(self):
        self.assertEqual(allocate_scores([1, 100], minimums=[20, 1], maximums=[100, 50]), [50, 50])
        self.assertEqual(allocate_scores([1, 100], minimums=[30, 1]), [30, 70])
        for lower, upper in (([60, 60], [100, 100]), ([1, 1], [40, 40]), ([True, 1], [100, 100])):
            with self.assertRaises(LessonValidationError) as error:
                allocate_scores([1, 1], minimums=lower, maximums=upper)
            self.assertEqual(error.exception.code, 'impossible_score_allocation')
        plan = weighted([1, 100])
        plan['assessment_items'][0]['points_min'] = 30
        result = normalize_meta_lesson(plan)
        self.assertEqual([i['points'] for i in result['assessment_items']], [30, 70])

    def test_task_counts_and_lesson_types_ru_kk(self):
        for language, code in (('Русский', 'ru'), ('Қазақша', 'kk')):
            for count in (1, 2, 3, 4, 5):
                for kind in ('Новый урок', 'Комбинированный урок', 'Практическое занятие', 'Зертханалық сабақ'):
                    with self.subTest(language=language, count=count, kind=kind):
                        plan = weighted(list(range(1, count + 1)), code)
                        result = normalize_meta_lesson(plan, language=language, lesson_type=kind)
                        self.assertEqual(len(result['assessment_items']), count)
                        self.assertEqual(sum(i['points'] for i in result['assessment_items']), 100)
                        self.assertEqual(sum(s['duration_minutes'] for s in result['lesson_stages']), 70)

    def test_descriptor_relative_weights_are_preserved(self):
        plan = ios()
        plan['assessment_items'][0]['descriptors'][0]['points'] = .5
        plan['assessment_items'][0]['descriptors'][1]['points'] = 1.5
        result = normalize_meta_lesson(plan)
        self.assertEqual(result['assessment_items'][0]['descriptor_points'], [5, 15])

    def test_missing_source_ids_are_assigned_deterministically_without_inventing_tasks(self):
        plan = ios()
        for stage in plan['lesson_stages']:
            for task in stage.get('tasks', []): task.pop('id')
        for item in plan['assessment_items']:
            item.pop('task_id')
            item['assessment_criterion'].pop('task_id')
            for descriptor in item['descriptors']: descriptor.pop('task_id')
        first = normalize_meta_lesson(plan)
        second = normalize_meta_lesson(plan)
        self.assertEqual([i['task_id'] for i in first['assessment_items']], [i['task_id'] for i in second['assessment_items']])
        self.assertEqual(sum(len(s['tasks']) for s in first['lesson_stages']), 5)

    def test_final_validator_rejects_unrendered_declared_activity(self):
        plan = normalize_meta_lesson(ios())
        stage = plan['lesson_stages'][1]
        stage['content'] = plan[stage['field_key']] = 'Другой материал без инструкции объявленного задания.'
        with self.assertRaises(LessonValidationError) as error: validate_meta_lesson(plan)
        self.assertEqual(error.exception.code, 'stage_task_content_mismatch')

    def test_conflicting_duplicate_descriptor_weights_are_rejected(self):
        plan = ios()
        plan['assessment_items'][0]['descriptor_points'] = [9, 11]
        with self.assertRaises(LessonValidationError) as error: normalize_meta_lesson(plan)
        self.assertEqual(error.exception.code, 'descriptor_points_mismatch')

    def test_final_validator_and_renderers_never_repair_canonical_scores(self):
        plan = normalize_meta_lesson(ios())
        original = copy.deepcopy(plan)
        validate_meta_lesson(plan)
        self.assertEqual(plan, original)
        plan['assessment_items'][0]['points'] = 19
        plan['assessment_items'][0]['descriptor_points'] = [9, 10]
        for operation in (lambda: validate_meta_lesson(plan), lambda: normalize_meta_lesson(plan),
                          lambda: _meta_lesson_body(plan, {'language': 'Русский'}),
                          lambda: build_doc('План урока', plan, template_path=get_college_config('META')['template'], college='META')):
            with self.assertRaises(LessonValidationError) as error: operation()
            self.assertEqual(error.exception.code, 'assessment_points_total')
            self.assertIn('99', str(error.exception))

    @patch('services.ai.openai_generate')
    def test_full_path_ai_json_preview_docx_atomic_supabase_same_model(self, provider):
        for language, code in (('Русский', 'ru'), ('Қазақша', 'kk')):
            source = json.loads((Path(__file__).with_name('fixtures') / 'meta_ios_95_ru.json').read_text()) if code == 'ru' else weighted([20, 20, 20, 20, 15], code)
            provider.return_value = '```json\n' + json.dumps(source, ensure_ascii=False) + '\n```'
            generated = create_lesson('Мобильная разработка', 'Операционная система iOS', language, 'IT', 'Chair', 'Новый урок', 70, None,
                                      college='META', lesson_date=date(2026, 10, 13))
            plan = json.loads(generated)
            client = Mock()
            client.rpc.side_effect = lambda name, params: Mock(execute=Mock(return_value=Mock(data=({'api_version': 2, 'ready': True, 'reasons': []} if name == 'lesson_plan_storage_status' else [{'lesson_id': 'fake', 'daily_count': 1}]))))
            service = SupabaseService.__new__(SupabaseService)
            service.client = client
            meta = replace(metadata(), language=language, topic='Операционная система iOS', lesson_type='Новый урок')
            doc, saved = save_generated_lesson_plan(service, generated, meta, get_college_config('META')['template'])
            stored = json.loads(client.rpc.call_args.args[1]['p_material_content'])
            self.assertEqual(plan, stored)
            self.assertEqual(client.rpc.call_count, 2)
            client.table.assert_not_called()
            provider.assert_called_once()
            provider.reset_mock()
            before = copy.deepcopy(plan)
            html = _meta_lesson_body(plan, {'language': language})
            html_scores = [int(v) for v in re.findall(r"<td style='[^']*'>(\d+)</td></tr>", html)]
            word = Document(doc)
            word_scores = [int(row.cells[4].text) for row in word.tables[2].rows[1:-1]]
            expected = [i['points'] for i in plan['assessment_items']]
            self.assertEqual(html_scores, expected)
            self.assertEqual(word_scores, expected)
            self.assertEqual(sum(word_scores), 100)
            self.assertEqual(len(word.tables[2].columns), 5)
            self.assertEqual(plan, before)
            stage_text = '\n'.join(c.text for row in word.tables[1].rows for c in row.cells)
            for stage in plan['lesson_stages']:
                for task in stage['tasks']:
                    self.assertIn(task['instruction'], stage_text)
            for item, row in zip(plan['assessment_items'], word.tables[2].rows[1:-1]):
                self.assertEqual(item['task_name'], row.cells[0].text)

    def test_diagnostics_include_actual_input_scores_and_anonymous_links(self):
        plan = weighted([20, 20, 20, 20, 15])
        snapshot = anonymized_snapshot(plan, LessonValidationError('bad'), 'normalization')
        self.assertEqual(snapshot['summary']['actual_points_total'], 95)
        self.assertEqual(snapshot['summary']['stage_count'], 7)
        self.assertEqual(snapshot['summary']['task_count'], 5)
        self.assertEqual([r['points'] for r in snapshot['summary']['points_by_task']], [20, 20, 20, 20, 15])
        self.assertEqual(len(snapshot['attempt_id']), 32)

