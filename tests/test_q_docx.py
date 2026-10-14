import json
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document

from config.colleges import get_college_config
from services.ai import create_lesson, AIServiceError
from services.docx_generator import build_doc
from services.lesson_validation import LessonValidationError
from services.q_docx import q_connections

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = get_college_config('Q')['template']
LINKS = {
    'Русский': 'Математика — повторяющиеся вычисления; логика — анализ условий и построение алгоритма',
    'Қазақша': 'Математика — қайталанатын есептеулер; логика — шарттарды талдау және алгоритм құру',
}


def q_fixture(language):
    return dict(topic='Python тіліндегі циклдер' if language == 'Қазақша' else 'Циклы Python',
                subject='Информатика', goal='Циклдерді қолдануды үйрету' if language == 'Қазақша' else 'Научить применять циклы',
                results='Алгоритм құрады' if language == 'Қазақша' else 'Создаёт алгоритм',
                interdisciplinary_connections=LINKS[language])


class QDocxTest(unittest.TestCase):
    def test_q_template_has_exactly_one_row_before_lesson_flow(self):
        original = Document(ROOT / 'docs/template_backups/template_q.original.docx')
        updated = Document(TEMPLATE)
        self.assertEqual(len(updated.tables), len(original.tables))
        self.assertEqual(len(updated.tables[0].rows), len(original.tables[0].rows) + 1)
        self.assertEqual(updated.tables[0].rows[-1].cells[0].text, 'Пәнаралық байланыс:')
        self.assertEqual(updated.tables[0].rows[-1].cells[1].text, '{interdisciplinary_connections}')
        self.assertIn('Ход урока', updated.paragraphs[6].text)

    def test_kazakh_and_russian_exports_are_localized_and_not_duplicated(self):
        for language, heading in (('Русский', 'Межпредметные связи:'), ('Қазақша', 'Пәнаралық байланыс:')):
            with self.subTest(language=language):
                file = build_doc('План урока', q_fixture(language), template_path=TEMPLATE,
                                 fields={'language': language, 'date': '2026-10-13'})
                doc = Document(file)
                matches = [r for t in doc.tables for r in t.rows if r.cells[0].text == heading]
                self.assertEqual(len(matches), 1)
                self.assertEqual(matches[0].cells[1].text, LINKS[language])
                self.assertIn('13.10.2026', doc.tables[0].rows[3].cells[1].text)
                self.assertEqual([r.cells[1].text for r in doc.tables[1].rows],
                                 [r.cells[1].text for r in Document(TEMPLATE).tables[1].rows])

    @patch('services.ai.openai_generate')
    def test_one_ai_request_includes_topic_specific_links_and_language(self, provider):
        provider.return_value = json.dumps(q_fixture('Қазақша'))
        plan = json.loads(create_lesson('Информатика', 'Python тіліндегі циклдер', 'Қазақша', 'IT',
                                       'Chair', 'Жаңа сабақ', 90, None, college='Q'))
        provider.assert_called_once()
        prompt = provider.call_args.args[0]
        for fragment in ('interdisciplinary_connections', '1–3', 'Python тіліндегі циклдер', 'Қазақша',
                         'знания или навыки', 'не утверждай его обязательность'):
            self.assertIn(fragment, prompt)
        self.assertEqual(plan['interdisciplinary_connections'], LINKS['Қазақша'])

    def test_subject_lists_without_explanations_are_rejected(self):
        for value in ('', 'Физика; математика', 'Физика —', '; '.join(['Математика — расчёт'] * 4)):
            with self.subTest(value=value), self.assertRaises(LessonValidationError):
                q_connections({'interdisciplinary_connections': value})

    def test_template_geometry_other_tables_and_package_parts_preserved(self):
        original = Document(ROOT / 'docs/template_backups/template_q.original.docx')
        updated = Document(TEMPLATE)
        self.assertEqual([p._p.xml for p in original.paragraphs], [p._p.xml for p in updated.paragraphs])
        for before, after in zip(original.tables, updated.tables):
            self.assertEqual([c.width for c in before.columns], [c.width for c in after.columns])
        for i in (1, 2):
            self.assertEqual(original.tables[i]._tbl.xml, updated.tables[i]._tbl.xml)
        self.assertEqual([r._tr.xml for r in original.tables[0].rows],
                         [r._tr.xml for r in updated.tables[0].rows[:-1]])
        self.assertEqual([s._sectPr.xml for s in original.sections], [s._sectPr.xml for s in updated.sections])
        with ZipFile(ROOT / 'docs/template_backups/template_q.original.docx') as before, ZipFile(TEMPLATE) as after:
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_other_colleges_never_get_the_new_q_row(self):
        for college in ('META', 'ETEC', 'ALT', 'ATKK'):
            doc = Document(get_college_config(college)['template'])
            self.assertFalse(any('{q_interdisciplinary_label}' in c.text or
                                 c.text.strip() == 'Пәнаралық байланыс:'
                                 for t in doc.tables for r in t.rows for c in r.cells))
