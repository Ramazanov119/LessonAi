"""Fill existing META tables while retaining their OOXML geometry and styling."""
from copy import deepcopy

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Twips, Pt

from services.meta_formatting import dynamic_run_font
from services.meta_localization import HEADERS, meta_labels

from services.lesson_validation import LessonValidationError, language_code



def set_cell_text(cell, text):
    """Retain cell/paragraph/first-run properties, replace only displayed text."""
    paragraph = cell.paragraphs[0]
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    for other in list(paragraph.runs)[1:]:
        other._r.getparent().remove(other._r)
    for other in list(cell.paragraphs)[1:]:
        other._p.getparent().remove(other._p)
    run.text = str(text)
    dynamic_run_font(run)
    # Content must wrap and rows must grow for long descriptors.
    for element in cell._tc.xpath('./w:tcPr/w:noWrap'):
        element.getparent().remove(element)


def _flexible_row(row):
    for element in row._tr.xpath('./w:trPr/w:trHeight'):
        element.getparent().remove(element)
    for paragraph in row._tr.xpath('.//w:pPr'):
        for name in ('keepNext', 'keepLines'):
            for element in paragraph.findall(qn('w:' + name)):
                paragraph.remove(element)


def find_meta_tables(doc):
    lesson = next((table for table in doc.tables if any(
        'Ход занятия' in cell.text or 'Сабақтың барысы' in cell.text
        for row in table.rows for cell in row.cells)), None)
    assessment = next((table for table in doc.tables if len(table.columns) == 5 and
                       any(word in ' '.join(cell.text for cell in table.rows[0].cells) for word in ('Блум', 'Bloom'))), None)
    if lesson is None or assessment is None:
        raise LessonValidationError('В META-шаблоне не найдены таблицы хода урока и оценивания.')
    return lesson, assessment


def fill_meta_tables(doc, payload, language):
    lesson, assessment = find_meta_tables(doc)
    lang = language_code(language)
    labels = meta_labels(language)
    first = next((i for i, row in enumerate(lesson.rows) if
                  '{org_steps}' in ' '.join(c.text for c in row.cells)), None)
    if first is None:
        raise LessonValidationError('В META-шаблоне не найден образец строки этапа.')
    prototype = deepcopy(lesson.rows[first]._tr)
    for row in list(lesson.rows)[first:]:
        lesson._tbl.remove(row._tr)
    resources = {
        'org_steps': 'presentation_resource', 'actualization': 'actualization_resource',
        'theory': 'presentation_demo', 'practice': 'practice_resource',
        'homework': 'homework_resource', 'assessment': 'assessment_sheet',
        'reflection': 'reflection_resource',
    }
    # Only the table's generated stage headers are localized, not the college shell.
    header_row = lesson.rows[first - 1]
    for cell, text in zip((header_row.cells[0], header_row.cells[1], header_row.cells[-1]), labels['stage_headers']):
        set_cell_text(cell, text)
    minute = labels['minute']
    for stage in payload['lesson_stages']:
        element = deepcopy(prototype)
        for merge in element.xpath('.//w:vMerge'):
            merge.getparent().remove(merge)
        lesson._tbl.append(element)
        row = _Row(element, lesson)
        _flexible_row(row)
        set_cell_text(row.cells[0], f"{stage['name']}\n({stage['duration_minutes']} {minute})")
        set_cell_text(row.cells[1], stage['content'])
        set_cell_text(row.cells[-1], payload.get(resources[stage['field_key']], '') or payload.get('resources', ''))

    prototype = deepcopy(assessment.rows[1]._tr)
    for row in list(assessment.rows)[1:]:
        assessment._tbl.remove(row._tr)
    for cell, text in zip(assessment.rows[0].cells, HEADERS[lang]):
        set_cell_text(cell, text)
    header_properties = assessment.rows[0]._tr.get_or_add_trPr()
    if header_properties.find(qn('w:tblHeader')) is None:
        header_properties.append(OxmlElement('w:tblHeader'))
    unit = labels['point']
    for item in payload['assessment_items']:
        element = deepcopy(prototype)
        assessment._tbl.append(element)
        row = _Row(element, assessment)
        _flexible_row(row)
        descriptors = ''
        values = (item['task_name'], item['bloom_level'], item['assessment_criterion'],
                  descriptors, str(item['points']))
        for cell, text in zip(row.cells, values):
            set_cell_text(cell, text)
        set_descriptors(row.cells[3], item['descriptors'], item['descriptor_points'], unit)
    element = deepcopy(prototype)
    assessment._tbl.append(element)
    row = _Row(element, assessment)
    _flexible_row(row)
    for cell, text in zip(row.cells, (labels['total'], '', '', '', '100')):
        set_cell_text(cell, text)
    format_assessment_table(doc, assessment)


def set_descriptors(cell, descriptions, scores, unit):
    # Separate paragraphs give each bullet its own hanging indent and spacing.
    paragraph = cell.paragraphs[0]
    paragraph_properties = deepcopy(paragraph._p.pPr)
    run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for index, (text, points) in enumerate(zip(descriptions, scores)):
        current = paragraph if index == 0 else cell.add_paragraph()
        if index:
            if paragraph_properties is not None:
                current._p.insert(0, deepcopy(paragraph_properties))
            run = current.add_run()
            if run_properties is not None:
                run._r.insert(0, deepcopy(run_properties))
        else:
            run = current.runs[0]
        run.text = f"• {text} — {points} {unit}"
        dynamic_run_font(run)
        current.paragraph_format.left_indent = Twips(120)
        current.paragraph_format.first_line_indent = Twips(-120)
        current.paragraph_format.space_after = Pt(3)


def format_assessment_table(doc, table):
    # This table is at the end of the document, in the last section.
    section = doc.sections[-1]
    available = (section.page_width - section.left_margin - section.right_margin) // Twips(1)
    # Task / Bloom / criterion / descriptors / points. Descriptors get 35%.
    proportions = (21, 16, 21, 35, 7)
    widths = [available * proportion // 100 for proportion in proportions]
    widths[3] += available - sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    width = properties.find(qn('w:tblW'))
    width.set(qn('w:type'), 'dxa')
    width.set(qn('w:w'), str(available))
    # Subtle rules separate the five columns without affecting other tables.
    borders = properties.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        properties.append(borders)
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
    indentation = properties.find(qn('w:tblInd'))
    if indentation is not None:
        indentation.set(qn('w:w'), '0')
    margins = properties.find(qn('w:tblCellMar'))
    if margins is None:
        margins = OxmlElement('w:tblCellMar')
        properties.append(margins)
    for side, padding in (('left', 80), ('right', 80), ('top', 60), ('bottom', 60)):
        edge = margins.find(qn('w:' + side))
        if edge is None:
            edge = OxmlElement('w:' + side)
            margins.append(edge)
        edge.set(qn('w:w'), str(padding))
        edge.set(qn('w:type'), 'dxa')
    spacing = properties.find(qn('w:tblCellSpacing'))
    if spacing is not None:
        spacing.set(qn('w:w'), '0')
    for column, value in zip(table.columns, widths):
        column.width = Twips(value)
    for row_index, row in enumerate(table.rows):
        _flexible_row(row)
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn('w:cantSplit')) is None:
            row_properties.append(OxmlElement('w:cantSplit'))
        for index, (cell, value) in enumerate(zip(row.cells, widths)):
            cell.width = Twips(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or index in (1, 4)
                                       else WD_ALIGN_PARAGRAPH.LEFT)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3 if index == 3 else 0)
                if index != 3 or row_index == 0:
                    paragraph.paragraph_format.first_line_indent = Twips(0)
                    paragraph.paragraph_format.left_indent = Twips(0)
                for run in paragraph.runs:
                    dynamic_run_font(run)
